#!/usr/bin/env python3
"""
generate.py - build Cisco device configurations from an Excel file.

Overview
--------
The script reads the link inventory in netgen_input.xlsx, turns it into one
data structure per device, and fills Jinja2 templates with it. Each selected
device gets one .cfg file.

    Excel  ->  read_sheet()  ->  build_devices()  ->  Jinja2 templates  ->  output/<device>.cfg
               (read + map       (check rows, build    (main.j2 includes
                the columns)      one "box" per device)  vrf/interfaces/bgp)

Usage
-----
    python generate.py [netgen_input.xlsx] [output_dir] [options]

    netgen_input.xlsx   Excel file to read (default: next to this script)
    output_dir          where the .cfg files are written (default: ./output)

    --templates FOLDER  use main.j2 from FOLDER for every device, instead of
                        templates/<os>/ chosen from the OS column of the
                        Generate sheet. Meant for drafts made by conf2jinja.py:
                        only what the Excel data provides is generated (loops
                        over lists the Excel doesn't fill produce nothing,
                        lines using a value the Excel doesn't have are left
                        out) and everything left out is listed.
    --strict            with --templates: stop on a missing value instead of
                        leaving it out.
    --debug             print the data given to the templates for each device.

Requirements: Python 3.8+, pip install openpyxl jinja2

Excel workbook
--------------
Sheet names and headers are matched without caring about case, spaces, "_" or
"-". Columns can be in any order, extra columns are ignored, and the header row
can be anywhere in the first 10 rows. Merged cells are read as if every cell of
the block had the value.

  Links     VRF | Subnet | DOT1Q | Vlan | DEVICE A | Port | IP | Device B | Port | IP
            One row per point-to-point link. BOTH ends get a configuration:
              DEVICE A: sub-interface Port A.DOT1Q, IP A, BGP neighbor IP B
              Device B: sub-interface Port B.DOT1Q, IP B, BGP neighbor IP A
            A plain "Port" / "IP" header belongs to the nearest DEVICE A /
            Device B column on its left; "Port A", "IP B"... work anywhere.
            Port B may stay empty when Device B is not generated.
            Required: VRF, Subnet, DOT1Q, DEVICE A, Port A, IP A, Device B, IP B.

  VRFs      VRF | VPN ID                 rd <local AS>:<VPN ID>

  Devices   Device | BGP AS [| Router ID]
            Every device of the Links sheet (both sides). Router ID is optional.

  Generate  Device | Generate | OS | VRF | Interfaces | BGP | Output name   (optional)
            Which devices to produce and how. yes/no columns: empty = yes.
            OS selects templates/<os>/ (only ios-xe exists for now).
            Output name: file name without .cfg (empty = hostname).
            Missing or empty sheet: every device, ios-xe, all blocks.

Data given to the templates (one per device)
--------------------------------------------
    local_as        "65001"
    router_id       "10.255.0.1"            only if filled in the Devices sheet
    blocks          {"vrf": True, "interfaces": True, "bgp": True}
    vrfs            [{name, vpn_id}]
    subinterfaces   [{name, description, dot1q, vrf, ip, mask}]
                      description = # <remote host> - <remote port>.<dot1q> - ( <VRF> ) #
    bgp_vrfs        [{vrf, vpn_id, neighbors: [{ip, remote_as, device_b, description}]}]
                      device_b = the device at the other end of the link
                      description = <remote host>-(vrf <VRF>)

Run with --debug to see this structure for your own data.

Exit codes: 0 = all configs written, 1 = nothing written or a device failed.

"""
import argparse
import ipaddress
import json
import re
import sys
from pathlib import Path

from jinja2 import ChainableUndefined, Environment, FileSystemLoader, StrictUndefined
from jinja2.exceptions import TemplateError, UndefinedError
from openpyxl import load_workbook

# Folder of this script: templates/ and the default Excel file are found from here,
# whatever folder the script is started from.
HERE = Path(__file__).resolve().parent

# Marker printed instead of a missing value (lenient mode, see SkipUndefined).
# Any rendered line containing it is removed afterwards. NUL characters cannot
# appear in a real config, so the marker can't collide with real text.
SKIP = "\x00SKIP\x00"


# =============================================================================
# Excel column definitions
# =============================================================================
# Each dictionary maps an internal name -> the header spellings accepted for it,
# written in normalized form (see norm()). To accept a new spelling used in
# another Excel file, add it here: nothing else needs to change.

LINKS_COLS = {
    "VRF": {"vrf"},
    "Subnet": {"subnet"},
    "DOT1Q": {"dot1q"},
    "Vlan": {"vlan"},                       # read but not used by the templates yet
    "DEVICE A": {"device a", "devicea"},
    "Device B": {"device b", "deviceb"},
    "Port_A": {"port a", "porta", "a port", "interface a"},
    "Port_B": {"port b", "portb", "b port", "interface b"},
    "IP_A": {"ip a", "ipa", "a ip"},
    "IP_B": {"ip b", "ipb", "b ip"},
}
LINKS_REQUIRED = ["VRF", "Subnet", "DOT1Q", "DEVICE A", "Device B", "Port_A", "IP_A", "IP_B"]

# Headers without a side ("Port", "IP"): they get the side (A or B) of the last
# DEVICE A / Device B column seen on their left. "Port" -> "Port_A" or "Port_B".
PLAIN = {"port": "Port", "interface": "Port", "ip": "IP"}

VRFS_COLS = {"VRF": {"vrf"}, "VPN ID": {"vpn id", "vpnid"}}

DEVICES_COLS = {
    "Device": {"device", "hostname"},
    "BGP AS": {"bgp as", "as", "asn", "bgp asn"},
    "Router ID": {"router id", "routerid"},
}

GEN_COLS = {
    "Device": {"device", "hostname"},
    "Generate": {"generate"},
    "OS": {"os", "os version", "platform"},
    "VRF": {"vrf"},
    "Interfaces": {"interfaces", "interface"},
    "BGP": {"bgp"},
    "Output name": {"output name", "output", "file name", "filename"},
}

# Template block -> its yes/no column in the Generate sheet.
# The block name is also the template file name: "bgp" -> bgp.j2
BLOCK_COLS = {"vrf": "VRF", "interfaces": "Interfaces", "bgp": "BGP"}

# Accepted yes/no values (normalized). An empty cell counts as yes.
YES = {"", "yes", "y", "x", "oui", "o", "1", "true"}
NO = {"no", "n", "non", "0", "false"}

# Accepted OS spellings (normalized) -> template folder name.
OS_NAMES = {"ios xe": "ios-xe", "iosxe": "ios-xe", "xe": "ios-xe",
            "ios xr": "ios-xr", "iosxr": "ios-xr", "xr": "ios-xr"}


# =============================================================================
# Jinja2 customizations
# =============================================================================

class SkipUndefined(ChainableUndefined):
    """Replacement for a value the Excel data doesn't provide (lenient mode).

    Jinja2 calls this class when a template uses a name that isn't in the data
    (for example {{ i.vlan }} or {% for l in loopbacks %}). Instead of failing:
      - looping over it gives nothing, so the loop body is never rendered;
      - printing it gives the SKIP marker, and that line is removed later.
    ChainableUndefined also makes deeper lookups (a.b.c) undefined instead of
    raising an error.

    Every missing name is stored in SkipUndefined.seen, so the script can tell
    you what was left out. main() empties it before each device.
    """
    seen = set()

    def _note(self):
        SkipUndefined.seen.add(self._undefined_name or "?")

    def __str__(self):
        self._note()
        return SKIP

    def __iter__(self):
        self._note()
        return iter(())


class BlockLoader(FileSystemLoader):
    """Template loader that serves switched-off blocks as empty files.

    When the Generate sheet says BGP = no, {% include 'bgp.j2' %} returns an
    empty template. main.j2 already checks `blocks`, but drafts from
    conf2jinja.py don't, so this makes the yes/no columns work everywhere.
    """

    def __init__(self, folder, disabled):
        """folder: template folder. disabled: block names, e.g. ["bgp"]."""
        super().__init__(str(folder))
        self.disabled = {f"{b}.j2" for b in disabled}

    def get_source(self, environment, template):
        if template in self.disabled:
            # (source, filename, uptodate): an empty template that never changes
            return "", None, lambda: True
        return super().get_source(environment, template)


# =============================================================================
# Reading the Excel file
# =============================================================================

class ExcelError(Exception):
    """A problem with the structure of the workbook (missing sheet or column,
    ambiguous header, invalid yes/no...). main() prints it and stops."""


def norm(text):
    """Normalize a header or value so spelling details don't matter.

    Lower case, trimmed, and runs of spaces / "_" / "-" become one space:
    'DEVICE A', 'Device_A', ' device-a ' -> 'device a'. None -> ''.
    """
    return re.sub(r"[\s_\-]+", " ", str(text or "").strip().lower())


def cell_text(value):
    """Turn an Excel cell value into clean text.

    None -> '' ; 2025.0 -> '2025' (Excel often stores whole numbers as floats);
    everything else -> str() without surrounding spaces.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def map_headers(raw, spec):
    """Match one row of header cells against a column definition.

    raw:  the cell values of the row, e.g. ['VRF', 'Subnet', ..., 'Port', 'IP']
    spec: a column definition such as LINKS_COLS
    Returns {column index: internal name}, e.g. {0: 'VRF', 5: 'Port_A', ...}.
    Columns not in spec are skipped. Raises ExcelError when a "Port"/"IP" header
    has no device column on its left, or when two columns mean the same thing.
    """
    mapping = {}
    side = None          # "A" or "B": the last device column seen, for PLAIN headers
    lookup = {alias: key for key, aliases in spec.items() for alias in aliases}

    for col, cell in enumerate(raw):
        h = norm(cell)
        if not h:
            continue
        key = lookup.get(h)

        if key in ("DEVICE A", "Device B"):
            side = "A" if key == "DEVICE A" else "B"

        # Plain "Port" / "IP": only meaningful in a sheet that has device sides
        if key is None and h in PLAIN and "DEVICE A" in spec:
            if side is None:
                raise ExcelError(f"column '{cell}' comes before any DEVICE A / Device B "
                                 f"column, so it's unclear which side it belongs to. "
                                 f"Rename it '{PLAIN[h]} A' or '{PLAIN[h]} B'.")
            key = f"{PLAIN[h]}_{side}"

        if key is None:
            continue                                  # extra column: ignored
        if key in mapping.values():
            raise ExcelError(f"two columns mean '{key.replace('_', ' ')}' "
                             f"(second one: '{cell}'). Rename one of them.")
        mapping[col] = key
    return mapping


def read_sheet(wb, name, spec, required):
    """Read one sheet into a list of row dictionaries.

    wb:       the workbook opened by openpyxl
    name:     sheet name to look for (case and separators don't matter)
    spec:     column definition (LINKS_COLS, VRFS_COLS...)
    required: internal names that must be present, e.g. LINKS_REQUIRED

    Every returned row has all keys of spec (empty string when the column is
    absent or the cell is empty) plus "_row", the real Excel row number used
    in error messages. Fully empty rows are skipped.
    Example: {"VRF": "underlay", "DOT1Q": "100", ..., "_row": 2}

    Raises ExcelError when the sheet or a required column can't be found.
    """
    ws = next((wb[s] for s in wb.sheetnames if norm(s) == norm(name)), None)
    if ws is None:
        raise ExcelError(f"sheet '{name}' not found (sheets: {', '.join(wb.sheetnames)})")

    # The whole sheet as a grid: rows[r][c], 0-based (Excel row 1 = rows[0])
    rows = [list(r) for r in ws.iter_rows(min_row=1, min_col=1, values_only=True)]

    # Merged cells: Excel keeps the value only in the top-left cell of the block,
    # the others read as empty. Copy it to every cell of the block, so a Port
    # merged over three links gives that Port to the three rows.
    for rng in ws.merged_cells.ranges:
        value = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, min(rng.max_row, len(rows)) + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if c <= len(rows[r - 1]):
                    rows[r - 1][c - 1] = value

    # Find the header row: among the first 10 rows, the one that matches the most
    # required columns. This allows title lines above the table.
    best, best_found, header_error = None, -1, None
    for i, raw in enumerate(rows[:10]):
        try:
            mapping = map_headers(raw, spec)
        except ExcelError as e:
            header_error = header_error or e          # keep the first one to report
            continue
        found = sum(1 for k in required if k in mapping.values())
        if found > best_found:
            best, best_found = (i, mapping), found
        if found == len(required):
            break
    if best is None or (best_found < len(required) and header_error):
        raise ExcelError(f"sheet '{ws.title}': {header_error or 'empty'}")

    header_row, mapping = best
    missing = [k.replace("_", " ") for k in required if k not in mapping.values()]
    if missing:
        raise ExcelError(f"sheet '{ws.title}': column(s) not found: {', '.join(missing)}. "
                         f"Headers seen: {[c for c in rows[header_row] if c]}")

    # Data rows -> dictionaries
    out = []
    for excel_row, raw in enumerate(rows[header_row + 1:], start=header_row + 2):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        row = {key: "" for key in spec}
        row["_row"] = excel_row
        row.update({key: cell_text(raw[col]) for col, key in mapping.items() if col < len(raw)})
        out.append(row)
    return out


def yes_no(value, where, column):
    """Read a yes/no cell. Empty counts as yes. Raises ExcelError otherwise.

    where/column are only used in the error message, e.g.
    "Generate row 3 (CE01): BGP must be yes or no, not 'maybe'".
    """
    v = norm(value)
    if v in YES:
        return True
    if v in NO:
        return False
    raise ExcelError(f"{where}: {column} must be yes or no, not '{value}'")


def read_generate(wb, hosts):
    """Read the Generate sheet and return the list of devices to produce.

    hosts: every device that has data (both sides of the links), used when the
           sheet is missing or empty.
    Returns (jobs, note):
      jobs: [{"host": "CE01", "os": "ios-xe",
              "blocks": {"vrf": True, "interfaces": True, "bgp": False},
              "output": "CE01"}]
      note: a message to print, or None.
    Rows with Generate = no are skipped. Raises ExcelError on invalid values.
    """
    everything = {b: True for b in BLOCK_COLS}

    def all_devices():
        return [{"host": h, "os": "ios-xe", "blocks": dict(everything), "output": h}
                for h in hosts]

    if not any(norm(s) == "generate" for s in wb.sheetnames):
        return all_devices(), None
    rows = read_sheet(wb, "Generate", GEN_COLS, ["Device"])
    if not rows:
        return all_devices(), "Generate sheet is empty: generating every device."

    jobs, seen = [], set()
    for r in rows:
        where = f"Generate row {r['_row']} ({r['Device']})"
        if not r["Device"]:
            continue
        if r["Device"] in seen:
            raise ExcelError(f"{where}: device listed twice")
        seen.add(r["Device"])
        if not yes_no(r["Generate"], where, "Generate"):
            continue

        os_name = OS_NAMES.get(norm(r["OS"]) or "ios xe")      # empty OS -> ios-xe
        if os_name is None:
            raise ExcelError(f"{where}: OS must be ios-xe or ios-xr, not '{r['OS']}'")

        blocks = {b: yes_no(r[col], where, col) for b, col in BLOCK_COLS.items()}
        if not any(blocks.values()):
            raise ExcelError(f"{where}: VRF, Interfaces and BGP are all 'no'")

        jobs.append({"host": r["Device"], "os": os_name, "blocks": blocks,
                     "output": r["Output name"] or r["Device"]})
    return jobs, None


# =============================================================================
# From link rows to one data structure per device
# =============================================================================

def build_devices(links, vpn_ids, devices):
    """Check every link row and build the template data of each device.

    links:   rows of the Links sheet (from read_sheet)
    vpn_ids: {"underlay": "100", ...}                        from the VRFs sheet
    devices: {"CE01": {"as": "65001", "router_id": ""}, ...} from the Devices sheet

    Each valid row is added twice, once for each end of the link (see add_side).
    Returns (data, errors, incomplete, no_desc):
      data:       {hostname: template data}, the structure described at the top
      errors:     list of messages; if not empty, nothing should be generated
      incomplete: {hostname: [row numbers]} for devices that are Device B on a
                  row with an empty Port B. They are removed from data, because
                  a config with only some of their links would be misleading.
      no_desc:    sub-interfaces without a description (Port B empty), e.g.
                  ["CE01 Gi0/0/1.100 (row 4)"]
    """
    errors, data, incomplete, no_desc = [], {}, {}, []
    used_subifs = {}        # (host, port, dot1q) -> where it was first defined

    def add_side(host, port, ip_local, ip_peer, peer, peer_port, vrf, dot1q, net, where):
        """Add one end of a link to the data of `host`.

        Called twice per row with the roles swapped:
          A side: host=DEVICE A, port=Port A, ip_local=IP A, ip_peer=IP B, peer=Device B
          B side: host=Device B, port=Port B, ip_local=IP B, ip_peer=IP A, peer=DEVICE A
        peer_port is the port at the other end, used in the interface description.
        """
        # The same sub-interface twice on one device is an Excel mistake
        key = (host, port, dot1q)
        if key in used_subifs:
            errors.append(f"{where}: {port}.{dot1q} already used on {host} "
                          f"({used_subifs[key]})")
            return
        used_subifs[key] = where

        # The device's data, created the first time the device is seen
        d = data.setdefault(host, {
            "local_as": devices[host]["as"],
            "vrfs": [], "subinterfaces": [], "bgp_vrfs": [],
        })
        if devices[host]["router_id"]:
            d["router_id"] = devices[host]["router_id"]

        # VRF definition: once per VRF on this device
        if not any(v["name"] == vrf for v in d["vrfs"]):
            d["vrfs"].append({"name": vrf, "vpn_id": vpn_ids[vrf]})

        # Sub-interface: one per link
        d["subinterfaces"].append({
            "name": port,
            # Standard: # <remote host> - <remote sub-interface> - ( <VRF> ) #
            # Empty (no description line) when the remote port is unknown.
            "description": (f"# {peer} - {peer_port}.{dot1q} - ( {vrf} ) #"
                            if peer_port else ""),
            "dot1q": dot1q,          # same value for <port>.<dot1q> and encapsulation
            "vrf": vrf,
            "ip": str(ip_local),
            "mask": str(net.netmask),
        })

        # BGP: one address-family per VRF, one neighbor per link
        af = next((a for a in d["bgp_vrfs"] if a["vrf"] == vrf), None)
        if af is None:
            af = {"vrf": vrf, "vpn_id": vpn_ids[vrf], "neighbors": []}
            d["bgp_vrfs"].append(af)
        af["neighbors"].append({
            "ip": str(ip_peer),
            "remote_as": devices[peer]["as"],
            "device_b": peer,                        # the device at the other end
            "description": f"{peer}-(vrf {vrf})",    # standard neighbor description
        })

    for r in links:
        where = f"Links row {r['_row']}"
        before = len(errors)          # to know if this row added errors
        host_a, host_b, vrf, dot1q = r["DEVICE A"], r["Device B"], r["VRF"], r["DOT1Q"]

        # --- checks: every problem is collected, so you see them all at once ---
        empty = [k.replace("_", " ") for k in LINKS_REQUIRED if not r[k]]
        if empty:
            errors.append(f"{where}: empty cell(s): {', '.join(empty)}")
            continue
        try:
            net = ipaddress.ip_network(r["Subnet"], strict=True)   # strict: no host bits
            ip_a = ipaddress.ip_address(r["IP_A"])
            ip_b = ipaddress.ip_address(r["IP_B"])
        except ValueError as e:
            errors.append(f"{where}: {e}")
            continue
        for label, ip in (("IP A", ip_a), ("IP B", ip_b)):
            if ip not in net:
                errors.append(f"{where}: {label} {ip} is not in {net}")
        if ip_a == ip_b:
            errors.append(f"{where}: IP A and IP B are both {ip_a}")
        if host_a == host_b:
            errors.append(f"{where}: DEVICE A and Device B are both {host_a}")
        for dev in (host_a, host_b):
            if dev not in devices:
                errors.append(f"{where}: device '{dev}' missing in Devices sheet")
        if vrf not in vpn_ids:
            errors.append(f"{where}: VRF '{vrf}' missing in VRFs sheet")
        if not dot1q.isdigit() or not 1 <= int(dot1q) <= 4094:
            errors.append(f"{where}: DOT1Q '{dot1q}' must be 1-4094")
        if len(errors) > before:
            continue

        # --- the row is valid: add both ends ---
        add_side(host_a, r["Port_A"], ip_a, ip_b, host_b, r["Port_B"], vrf, dot1q, net, where)
        if r["Port_B"]:
            add_side(host_b, r["Port_B"], ip_b, ip_a, host_a, r["Port_A"], vrf, dot1q, net, where)
        else:
            no_desc.append(f"{host_a} {r['Port_A']}.{dot1q} (row {r['_row']})")
            incomplete.setdefault(host_b, []).append(r["_row"])

    for host in incomplete:
        data.pop(host, None)
    return data, errors, incomplete, no_desc


# =============================================================================
# Engine: everything except printing and writing files
# =============================================================================
# generate_configs() does the work without printing or writing files; main()
# below handles the command line (messages, --debug, writing the .cfg files).

class Result:
    """Outcome of generate_configs().

    error_title / errors  set when nothing can be generated (Excel problems).
                          error_title is a sentence, errors the list of details.
    notes                 information to show, in order (warnings, defaults used).
    configs               one entry per generated device:
                            {"host", "os", "blocks", "output", "source",
                             "text": the configuration,
                             "data": what the templates received,
                             "left_out": names the Excel didn't provide}
    failed                devices whose templates failed: {"host", "lines": [...]}
    skipped               messages about devices not generated
    """

    def __init__(self):
        self.error_title = ""
        self.errors = []
        self.notes = []
        self.configs = []
        self.failed = []
        self.skipped = []

    @property
    def ok(self):
        """True when every selected device got its configuration."""
        return not self.errors and not self.failed

    def stop(self, title, errors):
        """Record a blocking problem and return self (used as `return res.stop(...)`)."""
        self.error_title, self.errors = title, list(errors)
        return self


def find_variant_loops(folder):
    """Names of the "variant" lists a draft loops over (bgp_vrfs_2, ...).

    conf2jinja.py creates one list per structure it found in the reference
    config (bgp_vrfs, bgp_vrfs_2...). Excel rows only fill the first one.
    """
    rx = r"{%-?\s*for\s+\w+\s+in\s+(?:\w+\.)?(\w+_\d+)\b"
    return sorted({m for f in folder.glob("*.j2")
                   for m in re.findall(rx, f.read_text(errors="replace"))})


def generate_configs(wb, templates=None, strict=False):
    """Build every selected configuration from an opened workbook.

    wb:        workbook from openpyxl.load_workbook(..., data_only=True)
    templates: optional folder with main.j2 used for every device (drafts);
               None = templates/<os>/ next to this script
    strict:    with `templates`, fail on a missing value instead of leaving it out

    Nothing is printed and no file is written: the caller decides what to do
    with the returned Result. Steps:
      1. read the Links, VRFs and Devices sheets
      2. build_devices(): check the rows, build each device's data
      3. read_generate(): which devices, OS and blocks
      4. check every job (data, templates, output names)
      5. render the templates
    Any problem in steps 1-4 returns a Result with no configs at all.
    """
    res = Result()

    def template_dir(os_name):
        """`templates` folder if given, else templates/<os>/."""
        return Path(templates).resolve() if templates else HERE / "templates" / os_name

    # --- 1. read the sheets ---------------------------------------------------
    try:
        links = read_sheet(wb, "Links", LINKS_COLS, LINKS_REQUIRED)
        vpn_ids = {r["VRF"]: r["VPN ID"]
                   for r in read_sheet(wb, "VRFs", VRFS_COLS, ["VRF", "VPN ID"])}
        devices = {r["Device"]: {"as": r["BGP AS"], "router_id": r["Router ID"]}
                   for r in read_sheet(wb, "Devices", DEVICES_COLS, ["Device", "BGP AS"])}
    except ExcelError as e:
        return res.stop("The Excel file has a structure problem:", [str(e)])

    # --- 2. build each device's data ------------------------------------------
    data, errors, incomplete, no_desc = build_devices(links, vpn_ids, devices)
    if errors:
        return res.stop("Fix these errors in the Excel file first:", errors)
    if not data and not incomplete:
        return res.stop("Nothing to generate:", ["No links found in the Links sheet."])
    if no_desc:
        res.notes.append("No interface description for these sub-interfaces, because "
                         "Port B (the facing interface) is empty: " + "; ".join(no_desc))

    # --- 3. what to generate --------------------------------------------------
    try:
        jobs, note = read_generate(wb, list(data))
    except ExcelError as e:
        return res.stop("The Generate sheet has a problem:", [str(e)])
    if note:
        res.notes.append(note)

    # --- 4. check every job before rendering anything -------------------------
    problems, outputs = [], {}
    for job in jobs:
        tpl_dir = template_dir(job["os"])
        if job["host"] in incomplete:
            rows = ", ".join(str(n) for n in incomplete[job["host"]])
            problems.append(f"{job['host']}: is Device B with an empty Port B "
                            f"(Links row(s) {rows})")
        elif job["host"] not in data:
            problems.append(f"{job['host']}: not found as DEVICE A or Device B in the Links sheet")
        elif not (tpl_dir / "main.j2").exists():
            what = f"in {tpl_dir}" if templates else f"for {job['os']}"
            problems.append(f"{job['host']}: no main.j2 {what} "
                            f"(expected {tpl_dir / 'main.j2'})")
        if job["output"] in outputs:
            problems.append(f"{job['host']}: output name '{job['output']}' already used "
                            f"by {outputs[job['output']]}")
        outputs[job["output"]] = job["host"]
    if problems:
        return res.stop("Nothing generated. Fix these first:", problems)
    if not jobs:
        return res.stop("Nothing to generate:",
                        ["No device has Generate = yes in the Generate sheet."])

    # Lenient mode: custom templates (drafts) rendered with Excel data only
    lenient = bool(templates) and not strict
    if lenient:
        variants = find_variant_loops(template_dir(""))
        if variants:
            res.notes.append(f"These loops are variants and stay empty: {', '.join(variants)}. "
                             "Excel rows are rendered with the structure of the first loop "
                             "(e.g. bgp_vrfs, not bgp_vrfs_2). Check the result, or merge "
                             "the variants in the draft.")

    # --- 5. render ------------------------------------------------------------
    for job in jobs:
        tpl_dir = template_dir(job["os"])
        disabled = [b for b, on in job["blocks"].items() if not on]
        env = Environment(
            loader=BlockLoader(tpl_dir, disabled),
            trim_blocks=True,                # no empty line left by {% for %}, {% if %}...
            keep_trailing_newline=True,      # keep the final newline of the files
            undefined=SkipUndefined if lenient else StrictUndefined,  # missing value: skip / error
        )
        d = data[job["host"]]
        SkipUndefined.seen = set()
        try:
            # **d passes each key as a template variable: d["local_as"] -> {{ local_as }}
            text = env.get_template("main.j2").render(blocks=job["blocks"], **d)
            if lenient:
                text = "\n".join(l for l in text.split("\n") if SKIP not in l)
        except UndefinedError as e:
            res.failed.append({"host": job["host"], "lines": [
                f"{job['host']}: the templates in {tpl_dir} use a variable the Excel "
                f"data doesn't provide: {e}",
                "  Available: local_as, router_id (if filled), blocks.vrf/interfaces/bgp,",
                "    vrfs[].name/vpn_id, subinterfaces[].name/dot1q/vrf/ip/mask/description,",
                "    bgp_vrfs[].vrf/vpn_id, bgp_vrfs[].neighbors[].ip/remote_as/device_b/description",
                "  Run without --strict to leave out what the Excel doesn't provide."]})
            continue
        except TemplateError as e:
            res.failed.append({"host": job["host"],
                               "lines": [f"{job['host']}: template error in {tpl_dir}: {e}"]})
            continue
        res.configs.append({
            "host": job["host"], "os": job["os"], "blocks": job["blocks"],
            "output": job["output"], "source": str(tpl_dir) if templates else job["os"],
            "text": text, "data": d,
            "left_out": sorted(SkipUndefined.seen) if lenient else [],
        })

    # --- devices that were not generated --------------------------------------
    selected = {j["host"] for j in jobs}
    skipped = [h for h in data if h not in selected]
    if skipped:
        res.skipped.append(f"Not selected in the Generate sheet: {', '.join(skipped)}")
    no_port = [h for h in incomplete if h not in selected]
    if no_port:
        res.skipped.append("Port B empty, so the Device B side can't be built: " +
                           ", ".join(f"{h} (row {', '.join(map(str, incomplete[h]))})"
                                     for h in no_port))
    return res


# =============================================================================
# Command line
# =============================================================================

def main():
    """Command-line entry point: options -> generate_configs() -> print and write.

    Nothing is written when the Excel file has a problem. Exit code 1 when
    nothing was written or a device failed.
    """
    ap = argparse.ArgumentParser(description="Excel -> device configs")
    ap.add_argument("xlsx", nargs="?", default=str(HERE / "netgen_input.xlsx"),
                    help="Excel file (default: netgen_input.xlsx next to this script)")
    ap.add_argument("outdir", nargs="?", default=str(HERE / "output"),
                    help="output folder (default: output/ next to this script)")
    ap.add_argument("--templates", help="template folder containing main.j2 "
                    "(default: templates/<os>/ next to this script)")
    ap.add_argument("--debug", action="store_true",
                    help="print the data given to the templates for each device")
    ap.add_argument("--strict", action="store_true",
                    help="with --templates: stop on a value the Excel doesn't provide "
                         "instead of leaving it out")
    args = ap.parse_args()

    xlsx, outdir = Path(args.xlsx), Path(args.outdir)
    if not xlsx.exists():
        print(f"Excel file not found: {xlsx.resolve()}")
        sys.exit(1)

    # data_only=True: read the results of Excel formulas, not the formulas
    wb = load_workbook(xlsx, data_only=True)
    res = generate_configs(wb, templates=args.templates, strict=args.strict)

    for note in res.notes:
        print(f"note: {note}")
    if res.errors:
        print(res.error_title)
        print("\n".join(f"  - {e}" for e in res.errors))
        sys.exit(1)

    outdir.mkdir(parents=True, exist_ok=True)
    for c in res.configs:
        if args.debug:
            # Exactly what the templates receive: {{ local_as }}, {% for i in subinterfaces %}...
            print(f"\n----- data for {c['host']} -----")
            print(json.dumps(c["data"], indent=2))
            print("-" * (len(c["host"]) + 21))
        path = outdir / f"{c['output']}.cfg"
        path.write_text(c["text"])
        parts = [b for b, on in c["blocks"].items() if on]
        print(f"written {path}  [{c['source']}: {', '.join(parts)}]")
        if c["left_out"]:
            print(f"  left out (not in Excel): {', '.join(c['left_out'])}")
    for f in res.failed:
        print("\n".join(f["lines"]))
    for s in res.skipped:
        print(f"not generated: {s}")
    if not res.ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
